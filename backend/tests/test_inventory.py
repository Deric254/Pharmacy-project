"""
Inventory tests. The properties that matter:
  1. Low-stock and expiry detection reflect real ledger state, not a
     stale cached number.
  2. Adjustments can never take a batch negative, and always require
     a reason - there's no path to a silent quantity edit.
  3. Reconciliation flags real mismatches without "fixing" anything
     itself (detection only).
  4. A sale that drops a product below its reorder point actually
     publishes stock.low - proving the event hook, not just the query.
"""

import asyncio
import json
from datetime import date, timedelta

from app.core.database import AsyncSessionLocal
from app.core.events import CHANNEL
from app.core.redis_client import redis_client
from app.models.medicine_batch import MedicineBatch
from app.models.product import Product
from app.models.stock_movement import MovementType, StockMovement
from app.models.stock_take import StockTake


async def _login(client, username: str, password: str) -> str:
    r = await client.post("/api/v1/auth/login", json={"username": username, "password": password})
    assert r.status_code == 200, r.text
    return str(r.json()["access_token"])


async def _make_product(name: str, reorder_point: int = 10) -> int:
    async with AsyncSessionLocal() as db:
        product = Product(name=name, reorder_point=reorder_point)
        db.add(product)
        await db.commit()
        return int(product.id)


async def _add_batch(
    product_id: int, qty: int, expiry: str = "2027-01-01", batch_number: str = "B1"
) -> int:
    async with AsyncSessionLocal() as db:
        batch = MedicineBatch(
            product_id=product_id,
            batch_number=batch_number,
            expiry_date=date.fromisoformat(expiry),
            qty_received=qty,
            qty_remaining=qty,
            cost_price=2.0,
            selling_price=5.0,
        )
        db.add(batch)
        await db.flush()
        db.add(
            StockMovement(
                batch_id=batch.id,
                movement_type=MovementType.PURCHASE,
                quantity_delta=qty,
                created_by_user_id=None,
            )
        )
        await db.commit()
        return int(batch.id)


class TestLowStock:
    async def test_product_below_reorder_point_is_flagged(self, client, owner_user):
        product_id = await _make_product("Amoxicillin", reorder_point=20)
        await _add_batch(product_id, qty=5)

        token = await _login(client, "lucy", "S3curePass!")
        r = await client.get(
            "/api/v1/inventory/low-stock", headers={"Authorization": f"Bearer {token}"}
        )
        assert r.status_code == 200
        flagged_ids = {p["product_id"] for p in r.json()}
        assert product_id in flagged_ids

    async def test_product_above_reorder_point_not_flagged(self, client, owner_user):
        product_id = await _make_product("Paracetamol", reorder_point=5)
        await _add_batch(product_id, qty=50)

        token = await _login(client, "lucy", "S3curePass!")
        r = await client.get(
            "/api/v1/inventory/low-stock", headers={"Authorization": f"Bearer {token}"}
        )
        flagged_ids = {p["product_id"] for p in r.json()}
        assert product_id not in flagged_ids

    async def test_product_with_zero_batches_counts_as_zero_stock(self, client, owner_user):
        product_id = await _make_product("Never Stocked", reorder_point=1)

        token = await _login(client, "lucy", "S3curePass!")
        r = await client.get(
            "/api/v1/inventory/low-stock", headers={"Authorization": f"Bearer {token}"}
        )
        flagged = {p["product_id"]: p for p in r.json()}
        assert product_id in flagged
        assert flagged[product_id]["total_qty_available"] == 0

    async def test_requires_permission(self, client, employee_user):
        # Employee fixture only has sales.create + inventory.view in
        # this suite's seeded_roles -- inventory.view IS granted, so
        # this should succeed; verifies view access works for the
        # role that actually needs to check stock while selling.
        token = await _login(client, "joe", "pass1234")
        r = await client.get(
            "/api/v1/inventory/low-stock", headers={"Authorization": f"Bearer {token}"}
        )
        assert r.status_code == 200

    async def test_results_are_ordered_lowest_stock_first(self, client, owner_user):
        # Three products, deliberately inserted in an order that does
        # NOT match ascending stock, so a passing test can only mean
        # the endpoint itself is sorting -- not that insertion order
        # happened to already look right.
        mid_id = await _make_product("Mid Stock Product", reorder_point=50)
        await _add_batch(mid_id, qty=30)
        lowest_id = await _make_product("Lowest Stock Product", reorder_point=50)
        await _add_batch(lowest_id, qty=2)
        highest_id = await _make_product("Highest Stock Product", reorder_point=50)
        await _add_batch(highest_id, qty=45)

        token = await _login(client, "lucy", "S3curePass!")
        r = await client.get(
            "/api/v1/inventory/low-stock", headers={"Authorization": f"Bearer {token}"}
        )
        assert r.status_code == 200
        ids_in_order = [p["product_id"] for p in r.json()]
        # All three are below their 50-unit reorder point, so all
        # three appear -- only their relative order is under test.
        relevant_order = [pid for pid in ids_in_order if pid in {lowest_id, mid_id, highest_id}]
        assert relevant_order == [lowest_id, mid_id, highest_id]


class TestExpiringBatches:
    async def test_batch_within_default_window_is_flagged(self, client, owner_user):
        product_id = await _make_product("Insulin")
        near_expiry = (date.today() + timedelta(days=10)).isoformat()
        await _add_batch(product_id, qty=20, expiry=near_expiry)

        token = await _login(client, "lucy", "S3curePass!")
        r = await client.get(
            "/api/v1/inventory/expiring", headers={"Authorization": f"Bearer {token}"}
        )
        assert r.status_code == 200
        product_ids = {b["product_id"] for b in r.json()}
        assert product_id in product_ids

    async def test_batch_far_in_future_not_flagged_by_default_window(self, client, owner_user):
        product_id = await _make_product("Vitamin C")
        far_expiry = (date.today() + timedelta(days=365)).isoformat()
        await _add_batch(product_id, qty=20, expiry=far_expiry)

        token = await _login(client, "lucy", "S3curePass!")
        r = await client.get(
            "/api/v1/inventory/expiring", headers={"Authorization": f"Bearer {token}"}
        )
        product_ids = {b["product_id"] for b in r.json()}
        assert product_id not in product_ids

    async def test_custom_window_overrides_default(self, client, owner_user):
        product_id = await _make_product("Vitamin D")
        expiry_in_200_days = (date.today() + timedelta(days=200)).isoformat()
        await _add_batch(product_id, qty=20, expiry=expiry_in_200_days)

        token = await _login(client, "lucy", "S3curePass!")
        headers = {"Authorization": f"Bearer {token}"}

        default_window = await client.get("/api/v1/inventory/expiring", headers=headers)
        assert product_id not in {b["product_id"] for b in default_window.json()}

        wide_window = await client.get(
            "/api/v1/inventory/expiring?within_days=365", headers=headers
        )
        assert product_id in {b["product_id"] for b in wide_window.json()}

    async def test_fully_depleted_batch_not_flagged_even_if_expiring(self, client, owner_user):
        product_id = await _make_product("Depleted Product")
        near_expiry = (date.today() + timedelta(days=5)).isoformat()
        await _add_batch(product_id, qty=0, expiry=near_expiry)

        token = await _login(client, "lucy", "S3curePass!")
        r = await client.get(
            "/api/v1/inventory/expiring", headers={"Authorization": f"Bearer {token}"}
        )
        product_ids = {b["product_id"] for b in r.json()}
        assert product_id not in product_ids


class TestValuation:
    async def test_valuation_matches_qty_times_cost(self, client, owner_user):
        product_id = await _make_product("Bandages")
        await _add_batch(product_id, qty=100)  # cost_price=2.0 in helper

        token = await _login(client, "lucy", "S3curePass!")
        r = await client.get(
            "/api/v1/inventory/valuation", headers={"Authorization": f"Bearer {token}"}
        )
        assert r.status_code == 200
        product_row = next(p for p in r.json()["by_product"] if p["product_id"] == product_id)
        assert product_row["qty_on_hand"] == 100
        assert product_row["value"] == 200.0  # 100 * 2.0
        assert r.json()["total_value"] >= 200.0

    async def test_expired_stock_does_not_inflate_valuation(self, client, owner_user):
        """
        The real bug this closes: a product could show real stock
        value and a real "available" quantity even when every unit
        was expired and therefore unsellable -- a genuine mismatch
        between what the numbers claimed and what a real sale attempt
        would actually do.
        """
        from datetime import date, timedelta

        yesterday = (date.today() - timedelta(days=1)).isoformat()
        product_id = await _make_product("Expired Only Product")
        await _add_batch(product_id, qty=100, expiry=yesterday)

        token = await _login(client, "lucy", "S3curePass!")
        r = await client.get(
            "/api/v1/inventory/valuation", headers={"Authorization": f"Bearer {token}"}
        )
        row = next((p for p in r.json()["by_product"] if p["product_id"] == product_id), None)
        # Either excluded entirely or present with zero value/qty --
        # never claiming expired-only stock has real, sellable value.
        if row is not None:
            assert row["qty_on_hand"] == 0
            assert row["value"] == 0.0


class TestAdjustments:
    async def test_adjustment_requires_permission(self, client, employee_user):
        product_id = await _make_product("Gauze")
        batch_id = await _add_batch(product_id, qty=50)

        token = await _login(client, "joe", "pass1234")
        r = await client.post(
            "/api/v1/inventory/adjustments",
            json={"batch_id": batch_id, "quantity_delta": -5, "reason": "DAMAGED"},
            headers={"Authorization": f"Bearer {token}"},
        )
        assert r.status_code == 403

    async def test_adjustment_writes_ledger_and_updates_qty(self, client, owner_user):
        product_id = await _make_product("Cotton Wool")
        batch_id = await _add_batch(product_id, qty=50)

        token = await _login(client, "lucy", "S3curePass!")
        r = await client.post(
            "/api/v1/inventory/adjustments",
            json={
                "batch_id": batch_id,
                "quantity_delta": -8,
                "reason": "DAMAGED",
                "notes": "Water damage in storage",
            },
            headers={"Authorization": f"Bearer {token}"},
        )
        assert r.status_code == 201
        assert r.json()["qty_remaining_after"] == 42

        async with AsyncSessionLocal() as db:
            from sqlalchemy import select

            batch_result = await db.execute(
                select(MedicineBatch).where(MedicineBatch.id == batch_id)
            )
            assert batch_result.scalar_one().qty_remaining == 42

            ledger_result = await db.execute(
                select(StockMovement).where(
                    StockMovement.batch_id == batch_id,
                    StockMovement.movement_type == MovementType.ADJUSTMENT,
                )
            )
            adjustment_rows = ledger_result.scalars().all()
            assert len(adjustment_rows) == 1
            assert adjustment_rows[0].quantity_delta == -8
            assert "DAMAGED" in adjustment_rows[0].reason
            assert "Water damage" in adjustment_rows[0].reason

    async def test_a_batch_held_by_a_stock_take_cannot_be_adjusted(self, client, owner_user):
        product_id = await _make_product("Gauze Rolls")
        batch_id = await _add_batch(product_id, qty=50)
        token = await _login(client, "lucy", "S3curePass!")
        headers = {"Authorization": f"Bearer {token}"}
        started = await client.post(
            "/api/v1/stock-takes", json={"product_ids": [product_id]}, headers=headers
        )
        assert started.status_code == 201

        r = await client.post(
            "/api/v1/inventory/adjustments",
            json={"batch_id": batch_id, "quantity_delta": -8, "reason": "DAMAGED"},
            headers=headers,
        )

        assert r.status_code == 409
        assert "locked by an open stock take" in r.json()["detail"]
        async with AsyncSessionLocal() as db:
            from sqlalchemy import select

            batch = (
                await db.execute(select(MedicineBatch).where(MedicineBatch.id == batch_id))
            ).scalar_one()
            ledger = (
                (
                    await db.execute(
                        select(StockMovement).where(
                            StockMovement.batch_id == batch_id,
                            StockMovement.movement_type == MovementType.ADJUSTMENT,
                        )
                    )
                )
                .scalars()
                .all()
            )
        assert batch.qty_remaining == 50
        assert ledger == []

    async def test_adjustment_cannot_take_batch_negative(self, client, owner_user):
        product_id = await _make_product("Syringes")
        batch_id = await _add_batch(product_id, qty=5)

        token = await _login(client, "lucy", "S3curePass!")
        r = await client.post(
            "/api/v1/inventory/adjustments",
            json={"batch_id": batch_id, "quantity_delta": -10, "reason": "THEFT_OR_LOSS"},
            headers={"Authorization": f"Bearer {token}"},
        )
        assert r.status_code == 400

        async with AsyncSessionLocal() as db:
            from sqlalchemy import select

            batch_result = await db.execute(
                select(MedicineBatch).where(MedicineBatch.id == batch_id)
            )
            assert batch_result.scalar_one().qty_remaining == 5  # unchanged

    async def test_zero_quantity_delta_rejected_by_schema(self, client, owner_user):
        product_id = await _make_product("Masks")
        batch_id = await _add_batch(product_id, qty=20)

        token = await _login(client, "lucy", "S3curePass!")
        r = await client.post(
            "/api/v1/inventory/adjustments",
            json={"batch_id": batch_id, "quantity_delta": 0, "reason": "MISCOUNT"},
            headers={"Authorization": f"Bearer {token}"},
        )
        assert r.status_code == 422

    async def test_positive_adjustment_can_correct_an_undercount(self, client, owner_user):
        product_id = await _make_product("Thermometers")
        batch_id = await _add_batch(product_id, qty=10)

        token = await _login(client, "lucy", "S3curePass!")
        r = await client.post(
            "/api/v1/inventory/adjustments",
            json={"batch_id": batch_id, "quantity_delta": 3, "reason": "MISCOUNT"},
            headers={"Authorization": f"Bearer {token}"},
        )
        assert r.status_code == 201
        assert r.json()["qty_remaining_after"] == 13

    async def test_two_concurrent_adjustments_both_apply(self, client, owner_user):
        """
        The actual bug this closes: adjustment used
        SELECT...FOR UPDATE, which SQLite silently drops entirely (the
        same false-safety pattern already found and fixed for stock
        decrements, PO transitions, refund restocks, stock-take
        closes, and loyalty points this session). Two adjustments to
        the same batch landing at once -- two managers independently
        correcting the same miscount, say -- could see the same
        starting quantity and the second commit would silently
        overwrite the first's correct result.
        """
        product_id = await _make_product("Concurrent Adjustment Product")
        batch_id = await _add_batch(product_id, qty=50)
        token = await _login(client, "lucy", "S3curePass!")
        headers = {"Authorization": f"Bearer {token}"}

        async def adjust(delta: int):
            return await client.post(
                "/api/v1/inventory/adjustments",
                json={"batch_id": batch_id, "quantity_delta": delta, "reason": "MISCOUNT"},
                headers=headers,
            )

        results = await asyncio.gather(adjust(5), adjust(3), return_exceptions=True)
        status_codes = [r.status_code for r in results if not isinstance(r, Exception)]
        assert status_codes.count(201) == 2  # both adjustments genuinely succeed

        async with AsyncSessionLocal() as db:
            batch = await db.get(MedicineBatch, batch_id)
            # Started at 50, +5 and +3 both must land -- not 53 or 55
            # (one adjustment silently lost), exactly 58.
            assert batch.qty_remaining == 58


class TestReconciliation:
    async def test_correctly_maintained_batch_has_no_issues(self, client, owner_user):
        product_id = await _make_product("Reconciled Product")
        await _add_batch(product_id, qty=30)

        token = await _login(client, "lucy", "S3curePass!")
        r = await client.get(
            "/api/v1/inventory/reconcile", headers={"Authorization": f"Bearer {token}"}
        )
        assert r.status_code == 200
        assert r.json() == []

    async def test_detects_a_deliberately_corrupted_batch(self, client, owner_user):
        product_id = await _make_product("Corrupted Product")
        batch_id = await _add_batch(product_id, qty=30)

        # Simulate a direct DB edit bypassing the application layer --
        # qty_remaining changed without a matching ledger row.
        async with AsyncSessionLocal() as db:
            from sqlalchemy import select

            batch_result = await db.execute(
                select(MedicineBatch).where(MedicineBatch.id == batch_id)
            )
            batch = batch_result.scalar_one()
            batch.qty_remaining = 999
            await db.commit()

        token = await _login(client, "lucy", "S3curePass!")
        r = await client.get(
            "/api/v1/inventory/reconcile", headers={"Authorization": f"Bearer {token}"}
        )
        issues = {i["batch_id"]: i for i in r.json()}
        assert batch_id in issues
        assert issues[batch_id]["qty_remaining"] == 999
        assert issues[batch_id]["ledger_sum"] == 30
        assert issues[batch_id]["discrepancy"] == 969
        # The real fix: a real product name and batch number, not a
        # cryptic ID someone would have to cross-reference by hand.
        assert issues[batch_id]["product_name"]
        assert issues[batch_id]["batch_number"]


class TestStockMovementHistory:
    """
    The real gap this closes: StockMovement has always been the
    complete, correct ledger of every quantity change (see its own
    docstring -- sales, refunds, adjustments, write-offs, corrections,
    and purchase receipts all write here), but nothing ever exposed it
    through the API. It existed and was accurate; nobody could see it.
    """

    async def test_a_purchase_receipt_appears_in_the_history(self, client, owner_user):
        product_id = await _make_product("Movement History Product")
        batch_id = await _add_batch(product_id, qty=40)

        token = await _login(client, "lucy", "S3curePass!")
        r = await client.get(
            "/api/v1/inventory/movements",
            params={"batch_id": batch_id},
            headers={"Authorization": f"Bearer {token}"},
        )
        assert r.status_code == 200
        body = r.json()
        assert body["total"] == 1
        entry = body["entries"][0]
        assert entry["movement_type"] == "PURCHASE"
        assert entry["quantity_delta"] == 40
        assert entry["product_name"] == "Movement History Product"
        assert entry["batch_id"] == batch_id

    async def test_filters_by_product_and_movement_type(self, client, owner_user):
        product_a = await _make_product("Filter Product A")
        product_b = await _make_product("Filter Product B")
        batch_a = await _add_batch(product_a, qty=50, batch_number="FA1")
        await _add_batch(product_b, qty=60, batch_number="FB1")

        async with AsyncSessionLocal() as db:
            db.add(
                StockMovement(
                    batch_id=batch_a,
                    movement_type=MovementType.ADJUSTMENT,
                    quantity_delta=-5,
                    reason="damaged",
                    created_by_user_id=None,
                )
            )
            await db.commit()

        token = await _login(client, "lucy", "S3curePass!")
        headers = {"Authorization": f"Bearer {token}"}

        by_product = await client.get(
            "/api/v1/inventory/movements", params={"product_id": product_a}, headers=headers
        )
        assert by_product.json()["total"] == 2  # the PURCHASE + the ADJUSTMENT
        assert all(e["product_id"] == product_a for e in by_product.json()["entries"])

        by_type = await client.get(
            "/api/v1/inventory/movements",
            params={"product_id": product_a, "movement_type": "ADJUSTMENT"},
            headers=headers,
        )
        assert by_type.json()["total"] == 1
        assert by_type.json()["entries"][0]["reason"] == "damaged"

    async def test_a_purchase_order_receipt_traces_back_through_the_history(
        self, client, owner_user
    ):
        """
        Proves the fix made to purchasing_service.py actually connects
        end to end: receiving stock via a real purchase order shows up
        here too, not just batches created directly in tests.
        """
        product_id = await _make_product("PO Traced Product")
        token = await _login(client, "lucy", "S3curePass!")
        headers = {"Authorization": f"Bearer {token}"}
        supplier = await client.post(
            "/api/v1/suppliers", json={"name": "Traced Supplier"}, headers=headers
        )
        po = await client.post(
            "/api/v1/purchase-orders/quick-purchase",
            json={
                "supplier_id": supplier.json()["id"],
                "lines": [
                    {
                        "product_id": product_id,
                        "quantity": 25,
                        "batch_number": "PO-TRACE-1",
                        "expiry_date": "2027-06-30",
                        "unit_cost": 4.0,
                        "selling_price": 8.0,
                    }
                ],
            },
            headers=headers,
        )
        assert po.status_code == 201, po.text

        r = await client.get(
            "/api/v1/inventory/movements", params={"product_id": product_id}, headers=headers
        )
        entries = r.json()["entries"]
        assert len(entries) == 1
        assert entries[0]["movement_type"] == "PURCHASE"
        assert entries[0]["quantity_delta"] == 25

    async def test_requires_inventory_adjust_permission(self, client, employee_user):
        # employee_user has inventory.view but not inventory.adjust --
        # matching reconcile's own permission tier, since this is
        # equally audit-sensitive detail, not routine stock browsing.
        token = await _login(client, "joe", "pass1234")
        r = await client.get(
            "/api/v1/inventory/movements", headers={"Authorization": f"Bearer {token}"}
        )
        assert r.status_code == 403

    async def test_excel_export_returns_every_matching_movement_not_one_page(
        self, client, owner_user
    ):
        """
        The real gap this closes: every other ledger-shaped view
        (Sales, Audit Logs, Customers, Suppliers, Products) supports
        export -- Stock Movements never did, despite being exactly the
        kind of record a real accounting or compliance review needs.
        Also confirms export returns everything matching the filter,
        not just one page -- created 3 movements, requested with
        limit=1 to prove the export ignores that pagination entirely.
        """
        product_id = await _make_product("Exportable Movement Product")
        await _add_batch(product_id, qty=10, batch_number="EXP-M1")
        await _add_batch(product_id, qty=20, batch_number="EXP-M2")
        await _add_batch(product_id, qty=30, batch_number="EXP-M3")

        token = await _login(client, "lucy", "S3curePass!")
        headers = {"Authorization": f"Bearer {token}"}

        r = await client.get(
            "/api/v1/inventory/movements",
            params={"product_id": product_id, "limit": 1, "export": "excel"},
            headers=headers,
        )
        assert r.status_code == 200
        assert (
            r.headers["content-type"]
            == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )

        import io
        import zipfile

        assert zipfile.is_zipfile(io.BytesIO(r.content))

        from openpyxl import load_workbook

        workbook = load_workbook(io.BytesIO(r.content))
        sheet = workbook.active
        assert sheet is not None
        data_rows = list(sheet.iter_rows(min_row=2, values_only=True))
        # All 3 purchase movements, despite limit=1 on the request --
        # export must never silently under-report to match the
        # on-screen page size.
        assert len(data_rows) == 3
        assert {row[3] for row in data_rows} == {"PURCHASE"}  # Type column
        assert {row[4] for row in data_rows} == {10, 20, 30}  # Quantity change column

    async def test_json_export_is_still_the_default(self, client, owner_user):
        product_id = await _make_product("JSON Default Movement Product")
        await _add_batch(product_id, qty=15)

        token = await _login(client, "lucy", "S3curePass!")
        r = await client.get(
            "/api/v1/inventory/movements",
            params={"product_id": product_id},
            headers={"Authorization": f"Bearer {token}"},
        )
        assert r.status_code == 200
        assert r.headers["content-type"].startswith("application/json")
        assert r.json()["total"] == 1


class TestSaleTriggeredLowStockEvent:
    async def test_sale_dropping_below_reorder_point_publishes_stock_low(
        self, client, employee_user
    ):
        product_id = await _make_product("Event Test Product", reorder_point=10)
        await _add_batch(product_id, qty=12)

        pubsub = redis_client.pubsub()
        await pubsub.subscribe(CHANNEL)
        await pubsub.get_message(timeout=1)  # discard the subscribe confirmation

        token = await _login(client, "joe", "pass1234")
        r = await client.post(
            "/api/v1/sales",
            json={
                "items": [{"product_id": product_id, "quantity": 5}],
                "payments": [{"method": "CASH", "amount": 25.0}],
            },
            headers={"Authorization": f"Bearer {token}"},
        )
        assert r.status_code == 201  # 12 - 5 = 7, below reorder_point of 10

        found_stock_low = False
        for _ in range(10):
            message = await pubsub.get_message(timeout=1)
            if message and message["type"] == "message":
                envelope = json.loads(message["data"])
                if envelope["event_type"] == "stock.low":
                    assert envelope["payload"]["product_id"] == product_id
                    assert envelope["payload"]["qty_remaining"] == 7
                    found_stock_low = True
                    break

        await pubsub.unsubscribe(CHANNEL)
        assert (
            found_stock_low
        ), "Expected a stock.low event after the sale dropped below reorder point"


class TestExpiredStockWriteOff:
    """
    The properties that matter:
      1. Only genuinely expired batches (server-verified against
         business_today(), never trusted from the caller) can be
         written off this way -- not "expiring soon", not anything
         still sellable.
      2. Every write-off is atomic, respects an active stock-take
         lock exactly like every other stock-mutating action in this
         app, and leaves the StockMovement ledger and qty_remaining
         in agreement (reconcile() must never flag it).
      3. Every write-off is audit logged.
      4. Bulk write-off only touches what's actually expired, and
         summarizes in one audit entry rather than one per batch.
    """

    async def test_writing_off_an_expired_batch_zeroes_it_and_logs_both_ledgers(
        self, client, owner_user
    ):
        token = await _login(client, "lucy", "S3curePass!")
        headers = {"Authorization": f"Bearer {token}"}
        product_id = await _make_product("Expired Write-off Product")
        batch_id = await _add_batch(product_id, qty=15, expiry="2020-01-01", batch_number="WO1")

        r = await client.post(
            f"/api/v1/inventory/batches/{batch_id}/write-off-expired", headers=headers
        )
        assert r.status_code == 200, r.text
        assert r.json() == {
            "batch_id": batch_id,
            "quantity_written_off": 15,
            "qty_remaining_after": 0,
        }

        async with AsyncSessionLocal() as db:
            batch = await db.get(MedicineBatch, batch_id)
            assert batch.qty_remaining == 0

        reconcile = await client.get("/api/v1/inventory/reconcile", headers=headers)
        assert reconcile.json() == [], "Ledger must still agree with qty_remaining after write-off"

        logs = await client.get(
            "/api/v1/audit-logs",
            params={"action": "batch.written_off_expired"},
            headers=headers,
        )
        matching = [e for e in logs.json()["entries"] if e["entity_id"] == str(batch_id)]
        assert len(matching) == 1
        assert matching[0]["old_value"] == "qty_remaining=15"
        assert matching[0]["new_value"] == "qty_remaining=0"

    async def test_a_batch_that_has_not_expired_yet_is_rejected(self, client, owner_user):
        token = await _login(client, "lucy", "S3curePass!")
        headers = {"Authorization": f"Bearer {token}"}
        product_id = await _make_product("Not Yet Expired Product")
        batch_id = await _add_batch(product_id, qty=10, expiry="2029-01-01", batch_number="WO2")

        r = await client.post(
            f"/api/v1/inventory/batches/{batch_id}/write-off-expired", headers=headers
        )
        assert r.status_code == 400, r.text

        async with AsyncSessionLocal() as db:
            batch = await db.get(MedicineBatch, batch_id)
            assert batch.qty_remaining == 10, "Rejected write-off must not touch quantity at all"

    async def test_a_batch_locked_by_an_open_stock_take_is_rejected(self, client, owner_user):
        token = await _login(client, "lucy", "S3curePass!")
        headers = {"Authorization": f"Bearer {token}"}
        product_id = await _make_product("Locked Expired Product")
        batch_id = await _add_batch(product_id, qty=8, expiry="2020-01-01", batch_number="WO3")

        async with AsyncSessionLocal() as db:
            stock_take = StockTake(initiated_by_user_id=owner_user.id)
            db.add(stock_take)
            await db.flush()
            batch = await db.get(MedicineBatch, batch_id)
            batch.locked_by_stock_take_id = stock_take.id
            await db.commit()

        r = await client.post(
            f"/api/v1/inventory/batches/{batch_id}/write-off-expired", headers=headers
        )
        assert r.status_code == 409, r.text

        async with AsyncSessionLocal() as db:
            batch = await db.get(MedicineBatch, batch_id)
            assert batch.qty_remaining == 8

    async def test_a_batch_already_at_zero_is_rejected(self, client, owner_user):
        token = await _login(client, "lucy", "S3curePass!")
        headers = {"Authorization": f"Bearer {token}"}
        product_id = await _make_product("Already Zero Product")
        batch_id = await _add_batch(product_id, qty=5, expiry="2020-01-01", batch_number="WO4")

        first = await client.post(
            f"/api/v1/inventory/batches/{batch_id}/write-off-expired", headers=headers
        )
        assert first.status_code == 200, first.text

        second = await client.post(
            f"/api/v1/inventory/batches/{batch_id}/write-off-expired", headers=headers
        )
        assert second.status_code == 400, second.text

    async def test_employee_without_permission_is_forbidden(
        self, client, owner_user, employee_user
    ):
        product_id = await _make_product("Permission Test Product")
        batch_id = await _add_batch(product_id, qty=5, expiry="2020-01-01", batch_number="WO5")

        token = await _login(client, "joe", "pass1234")
        r = await client.post(
            f"/api/v1/inventory/batches/{batch_id}/write-off-expired",
            headers={"Authorization": f"Bearer {token}"},
        )
        assert r.status_code == 403

    async def test_bulk_write_off_only_touches_genuinely_expired_batches(self, client, owner_user):
        token = await _login(client, "lucy", "S3curePass!")
        headers = {"Authorization": f"Bearer {token}"}
        product_id = await _make_product("Bulk Write-off Product")
        expired_1 = await _add_batch(product_id, qty=10, expiry="2020-01-01", batch_number="BULK1")
        expired_2 = await _add_batch(product_id, qty=6, expiry="2021-06-15", batch_number="BULK2")
        still_good = await _add_batch(product_id, qty=20, expiry="2029-01-01", batch_number="BULK3")

        r = await client.post("/api/v1/inventory/write-off-all-expired", headers=headers)
        assert r.status_code == 200, r.text
        body = r.json()
        assert body["batches_written_off"] == 2
        assert body["total_quantity_written_off"] == 16
        assert {d["batch_id"] for d in body["details"]} == {expired_1, expired_2}

        async with AsyncSessionLocal() as db:
            assert (await db.get(MedicineBatch, expired_1)).qty_remaining == 0
            assert (await db.get(MedicineBatch, expired_2)).qty_remaining == 0
            assert (await db.get(MedicineBatch, still_good)).qty_remaining == 20

        reconcile = await client.get("/api/v1/inventory/reconcile", headers=headers)
        assert reconcile.json() == []

        logs = await client.get(
            "/api/v1/audit-logs",
            params={"action": "batch.bulk_written_off_expired"},
            headers=headers,
        )
        entries = logs.json()["entries"]
        assert len(entries) == 1
        assert "batches=2" in entries[0]["new_value"]
        assert "total_quantity=16" in entries[0]["new_value"]

    async def test_bulk_write_off_with_nothing_expired_is_a_clean_no_op(self, client, owner_user):
        token = await _login(client, "lucy", "S3curePass!")
        headers = {"Authorization": f"Bearer {token}"}
        product_id = await _make_product("Nothing Expired Product")
        await _add_batch(product_id, qty=10, expiry="2029-01-01", batch_number="NOEXP")

        r = await client.post("/api/v1/inventory/write-off-all-expired", headers=headers)
        assert r.status_code == 200, r.text
        assert r.json() == {
            "batches_written_off": 0,
            "total_quantity_written_off": 0,
            "details": [],
        }

    async def test_two_concurrent_write_offs_of_the_same_batch_never_double_log(
        self, client, owner_user
    ):
        """
        Proves the compare-and-swap guard, not just asserts it. Two
        requests racing to write off the exact same batch must result
        in exactly one success (quantity_written_off == the real
        original amount) and one clean rejection -- never two
        "successes" that would double-count the same stock as removed
        twice in the ledger.
        """
        token = await _login(client, "lucy", "S3curePass!")
        headers = {"Authorization": f"Bearer {token}"}
        product_id = await _make_product("Concurrent Write-off Product")
        batch_id = await _add_batch(product_id, qty=12, expiry="2020-01-01", batch_number="RACE1")

        async def write_off():
            return await client.post(
                f"/api/v1/inventory/batches/{batch_id}/write-off-expired", headers=headers
            )

        results = await asyncio.gather(write_off(), write_off(), return_exceptions=True)
        exceptions = [r for r in results if isinstance(r, Exception)]
        assert not exceptions, f"Unhandled exceptions under concurrency: {exceptions}"

        status_codes = sorted(r.status_code for r in results)
        # The loser already passed its own "not yet zero" pre-check
        # before the race resolved, so it lands on the compare-and-
        # swap failure branch (409: "changed just now, please retry")
        # rather than the early-bailout branch a genuinely later,
        # non-racing request would hit (400: "already zero") -- both
        # are correct rejections, this is just which one a true race
        # actually produces.
        assert status_codes == [200, 409], status_codes

        succeeded = next(r for r in results if r.status_code == 200)
        assert succeeded.json()["quantity_written_off"] == 12

        async with AsyncSessionLocal() as db:
            batch = await db.get(MedicineBatch, batch_id)
            assert batch.qty_remaining == 0

        reconcile = await client.get("/api/v1/inventory/reconcile", headers=headers)
        assert reconcile.json() == [], "Two racing write-offs must never double-count in the ledger"
