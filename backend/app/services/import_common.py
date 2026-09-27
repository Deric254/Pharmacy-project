"""
Shared helpers for bulk-import services (customer, product, purchase
order, stock take).

Every bulk import in this system builds its Excel template the same
way (a bold white-on-navy header row, an optional greyed-out example
row, a small red instructions note) and validates rows with the same
two conventions: track "have I seen this key before in this file?"
for duplicate detection, and gate the whole import behind a single
422 that lists every problem at once rather than importing whatever
happened to be clean. Pulling these out means a decision like "header
row is bold white on #1F2937" only has to be made, and fixed, once --
and the four importers can't quietly drift out of sync with each
other on wording or styling the way they had (see the duplicate-row
message: three of the four importers already agreed on "in this same
file.", one said "in this file." -- now there is only one wording to
agree with).
"""

from typing import Any

from fastapi import HTTPException
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet

from app.schemas.product import ImportRowError

HEADER_FONT = Font(name="Arial", bold=True, color="FFFFFF")
HEADER_FILL = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
EXAMPLE_FONT = Font(name="Arial", italic=True, color="6B7280")
INSTRUCTIONS_FONT = Font(name="Arial", italic=True, size=9, color="991B1B")


def clean_str(value: Any) -> str:
    """Trim a raw spreadsheet cell to a string; a blank/None cell becomes ''."""
    return str(value).strip() if value is not None else ""


def write_header_row(ws: Worksheet, headers: list[str]) -> None:
    """Write the standard bold white-on-navy header row and freeze it in place."""
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
    ws.freeze_panes = "A2"


def write_example_row(ws: Worksheet, example_row: list[Any]) -> None:
    """Write the standard greyed-out italic example row directly under the header."""
    for col, value in enumerate(example_row, start=1):
        cell = ws.cell(row=2, column=col, value=value)
        cell.font = EXAMPLE_FONT


def write_instructions(ws: Worksheet, row: int, column: int, text: str) -> None:
    """Write the small red italic instructions note next to the header row."""
    cell = ws.cell(row=row, column=column, value=text)
    cell.font = INSTRUCTIONS_FONT


def check_duplicate_in_file(
    seen: dict[str, int], key: str, row_num: int, field: str, errors: list[ImportRowError]
) -> bool:
    """
    Record the first row `key` was seen on in `seen`, appending a
    duplicate-row error to `errors` if it already appeared earlier in
    this same file. Returns True if this row was a duplicate, so a
    caller that needs to also mark the row invalid can do so -- not
    every importer treats an in-file duplicate as fatal to that row on
    its own (some rely on the all-or-nothing gate below to reject the
    file either way), so the decision is left to the caller.
    """
    if key in seen:
        errors.append(
            ImportRowError(
                row=row_num,
                field=field,
                message=f"Duplicate of row {seen[key]} in this same file.",
            )
        )
        return True
    seen[key] = row_num
    return False


def raise_if_errors(errors: list[ImportRowError], not_imported_word: str) -> None:
    """
    The all-or-nothing gate every bulk import shares: if anything at
    all is wrong anywhere in the file, raise one 422 listing every
    problem at once and import nothing -- never a partial import.
    """
    if errors:
        raise HTTPException(
            status_code=422,
            detail={
                "message": f"{len(errors)} problem(s) found. Nothing was {not_imported_word}.",
                "errors": [e.model_dump() for e in errors],
            },
        )
